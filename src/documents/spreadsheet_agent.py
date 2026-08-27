from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

class SpreadsheetAgent:
    def create(self, title, headers, rows, output_path):
        path=Path(output_path)
        path.parent.mkdir(parents=True,exist_ok=True)

        wb=Workbook()
        ws=wb.active
        ws.title="Dados"

        ws["A1"]=title or "Planilha"
        ws["A1"].font=Font(bold=True,size=16)

        start=3
        thin=Side(style="thin")
        for col,h in enumerate(headers or [],1):
            c=ws.cell(start,col,h)
            c.font=Font(bold=True)
            c.alignment=Alignment(horizontal="center")
            c.border=Border(bottom=thin)

        for r_idx,row in enumerate(rows or [],start+1):
            for c_idx,val in enumerate(row,1):
                ws.cell(r_idx,c_idx,val)

        for col in range(1,max(1,len(headers or []))+1):
            width=12
            for cell in ws[get_column_letter(col)]:
                if cell.value is not None:
                    width=max(width,min(40,len(str(cell.value))+2))
            ws.column_dimensions[get_column_letter(col)].width=width

        wb.save(path)
        return str(path)
