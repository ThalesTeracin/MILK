"""
Geracao de .pptx via python-pptx, com carregamento tardio.

Diferente de docx_agent, este agente ainda depende de python-pptx (e
portanto de lxml): gerar OOXML de apresentacao a mao exige slideMaster,
slideLayouts, tema e relacionamentos por slide, e um erro em qualquer um
faz o PowerPoint recusar o arquivo. Nao vale o risco enquanto DOCX, PDF e
XLSX cobrem o uso real.

O que muda aqui e *quando* a biblioteca e importada. Com o import no topo
do modulo, um lxml bloqueado pelo Smart App Control derrubava tambem o
import de document_manager, e com ele os tres formatos que funcionam. Com
o import dentro de _carregar(), a indisponibilidade fica contida neste
agente e vira erro tratado no gerente.
"""

from pathlib import Path


def _carregar():
    """
    Importa python-pptx sob demanda.

    Levanta ImportError quando a biblioteca esta ausente ou quando uma
    dependencia nativa dela foi bloqueada pelo sistema. Quem chama deve
    tratar e degradar; ver DocumentManager.create_from_request.
    """
    from pptx import Presentation
    from pptx.util import Pt

    return Presentation, Pt


class PptxAgent:
    def create(self, title, slides, output_path):
        Presentation, Pt = _carregar()

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        prs = Presentation()

        # Slide de titulo.
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title or "Apresentação"
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = "Gerado pela MILK"

        for item in slides or []:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = item.get("title", "Slide")
            tf = slide.placeholders[1].text_frame
            tf.clear()
            for i, line in enumerate(item.get("bullets", [])):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = str(line)
                p.font.size = Pt(20)

        prs.save(path)
        return str(path)
