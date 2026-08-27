from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

class DocxAgent:
    def create(self, title, body, output_path):
        path=Path(output_path)
        path.parent.mkdir(parents=True,exist_ok=True)

        doc=Document()
        styles=doc.styles
        styles["Normal"].font.name="Aptos"
        styles["Normal"].font.size=Pt(11)

        p=doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(title or "Documento")
        r.bold=True
        r.font.size=Pt(18)

        for block in (body or "").split("\n\n"):
            block=block.strip()
            if not block:
                continue
            if block.startswith("# "):
                doc.add_heading(block[2:].strip(),level=1)
            elif block.startswith("## "):
                doc.add_heading(block[3:].strip(),level=2)
            else:
                doc.add_paragraph(block)

        doc.save(path)
        return str(path)
